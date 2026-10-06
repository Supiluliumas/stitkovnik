"""Optional integration check. Start app.py --no-browser before running."""
from io import BytesIO
from pathlib import Path
import sys

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parent.parent
URL = sys.argv[1] if len(sys.argv) > 1 else 'http://127.0.0.1:8765'


def main():
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        context = browser.new_context(viewport={"width": 1440, "height": 1080}, accept_downloads=True)
        page = context.new_page()
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(URL)
        expect(page.locator('#countBadge')).to_have_text('0')
        page.screenshot(path='/tmp/stitkovnik-empty.png', full_page=True)
        page.locator('#filesInput').set_input_files([str(p) for p in sorted((ROOT / 'ukazkove-stitky').glob('*.png'))])
        expect(page.locator('#countBadge')).to_have_text('2', timeout=120000)
        expect(page.locator('#progressSection')).to_be_hidden(timeout=10000)
        expect(page.get_by_role('textbox', name='S/N: kamera-01.png', exact=True)).to_have_value('000123456789')
        expect(page.get_by_role('textbox', name='MAC: kamera-02-otocena.png', exact=True)).to_have_value('00:11:22:33:44:55')
        # A repeated identifier flags both rows, and disappears after correction.
        serial = page.get_by_role('textbox', name='S/N: kamera-02-otocena.png', exact=True)
        serial.fill('000123456789')
        expect(page.locator('#reviewCount')).to_have_text('2')
        serial.fill('000987654321')
        expect(page.locator('#reviewCount')).to_have_text('0')
        page.get_by_role('button', name='kamera-01.png', exact=True).click()
        expect(page.locator('#preview')).to_be_visible()
        page.locator('#detailFields input[data-key="Model"]').fill('DS-2CD2043G2-I')
        page.locator('#verified').check()
        page.locator('#saveDetailButton').click()
        expect(page.locator('#readyCount')).to_have_text('1')
        # Custom label mapping and reparsing of a previously scanned image.
        page.locator('#rulesButton').click()
        page.once('dialog', lambda dialog: dialog.accept('Umístění'))
        page.locator('#newRuleButton').click()
        page.locator('#rulesFields input[data-key="Umístění"]').fill('Location, Umístění')
        page.locator('#saveRulesButton').click()
        expect(page.locator('#rulesDialog')).not_to_be_visible()
        page.get_by_role('button', name='kamera-02-otocena.png', exact=True).click()
        raw = page.locator('#rawText')
        raw.fill(raw.input_value() + '\nLocation: Brána A')
        page.locator('#reparseButton').click()
        expect(page.locator('#detailFields input[data-key="Umístění"]')).to_have_value('Brána A')
        page.locator('#saveDetailButton').click()
        # Reload must preserve text values and verification without persisting photos.
        page.reload()
        expect(page.locator('#countBadge')).to_have_text('2')
        expect(page.locator('#readyCount')).to_have_text('1')
        page.locator('#search').fill('000123456789')
        expect(page.locator('#rows tr')).to_have_count(1)
        with page.expect_download() as download_event:
            page.locator('#exportButton').click()
        download = download_event.value
        workbook = load_workbook(BytesIO(Path(download.path()).read_bytes()))
        sheet = workbook.active
        assert sheet.max_row == 3, 'Export must include filtered-out rows'
        assert sheet['B2'].value == '000123456789' and sheet['B2'].data_type == 's'
        assert sheet['C3'].value == '00:11:22:33:44:55'
        headers = [cell.value for cell in sheet[1]]
        assert sheet.cell(3, headers.index('Umístění') + 1).value == 'Brána A'
        with page.expect_download() as csv_event:
            page.locator('#csvButton').click()
        assert Path(csv_event.value.path()).read_bytes().startswith(b'\xef\xbb\xbf')
        page.locator('#search').fill('')
        page.screenshot(path='/tmp/stitkovnik-results.png', full_page=True)
        # Check responsive layout without horizontal page overflow.
        page.set_viewport_size({"width": 390, "height": 844})
        page.screenshot(path='/tmp/stitkovnik-mobile.png', full_page=True)
        assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
        page.set_viewport_size({"width": 1440, "height": 1080})
        page.get_by_role('button', name='kamera-02-otocena.png', exact=True).click()
        page.locator('#detailNote').fill('Poznámka musí přežít nové OCR')
        page.locator('#saveDetailButton').click()
        page.get_by_role('textbox', name='S/N: kamera-02-otocena.png', exact=True).fill('BROKEN123')
        page.locator('#filesInput').set_input_files([str(p) for p in sorted((ROOT / 'ukazkove-stitky').glob('*.png'))])
        expect(page.locator('#progressSection')).to_be_hidden(timeout=120000)
        expect(page.locator('#countBadge')).to_have_text('2')
        expect(page.locator('#readyCount')).to_have_text('1')
        expect(page.get_by_role('textbox', name='S/N: kamera-02-otocena.png', exact=True)).to_have_value('000987654321')
        page.get_by_role('button', name='kamera-02-otocena.png', exact=True).click()
        expect(page.locator('#detailNote')).to_have_value('Poznámka musí přežít nové OCR')
        page.locator('[data-close="detailDialog"]').first.click()
        page.locator('#clearButton').click()
        page.locator('#confirmClear').click()
        expect(page.locator('#countBadge')).to_have_text('0')
        assert not errors, errors
        # Write endpoints reject unrelated pages without the session token.
        response = context.request.post(URL + '/api/reparse', data='{}')
        assert response.status == 403
        browser.close()
    print('Browser check passed: OCR, rotation, duplicates, corrections, rules, reload, XLSX/CSV, mobile, reprocessing, preserved verification/notes, clear, access checks.')


if __name__ == '__main__':
    main()
